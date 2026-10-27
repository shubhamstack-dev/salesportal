"""Small Excel writer: {sheet name: [row dicts]} -> .xlsx bytes."""
import io

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def workbook(sheets: dict[str, list[dict]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        heads = list(rows[0].keys()) if rows else ["No rows"]
        ws.append(heads)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="14213D")
        for r in rows:
            ws.append([r.get(h) for h in heads])
        for i, h in enumerate(heads, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(str(h)) + 3))
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(data: bytes, filename: str) -> Response:
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})
