"""End-to-end API tests. SQLite by default; set TEST_DATABASE_URL for MySQL.
The master data comes from tests/Sample_Dynamics_Masters.xlsx, as it would from Dynamics."""
import io
import os
import tempfile
from datetime import timedelta
from pathlib import Path

os.environ.setdefault("DATABASE_URL", os.getenv("TEST_DATABASE_URL", f"sqlite:///{tempfile.mkdtemp()}/t.db"))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Attendance, User  # noqa: E402
from app.services.common import local_at, local_today  # noqa: E402

SAMPLE = Path(__file__).parent / "Sample_Dynamics_Masters.xlsx"
LOC = {"lat": 26.8470, "lng": 80.9460, "acc": 12}


@pytest.fixture(scope="module")
def c():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as client:
        yield client


def H(t):
    return {"Authorization": f"Bearer {t}"}


def login(c, u, p):
    r = c.post("/api/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def admin(c):
    assert c.get("/api/auth/status").json()["needs_setup"]
    r = c.post("/api/auth/setup", json={"name": "Randeep Kapoor", "username": "randeep", "password": "admin12345"})
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def masters(c, admin):
    for ent in ["customers", "materials", "attributes", "prices", "hsn", "charges"]:
        r = c.post(f"/api/integration/import/{ent}/file", files={"file": ("masters.xlsx", SAMPLE.read_bytes())}, headers=H(admin))
        assert r.status_code == 200, r.text
        p = r.json()
        assert not p["result"]["errors"], (ent, p["result"])
        r = c.post(f"/api/integration/import/{ent}/apply", json={"import_id": p["import_id"], "sheet": p["sheet"], "map": p["map"]}, headers=H(admin))
        assert r.status_code == 200, r.text
    return True


@pytest.fixture(scope="module")
def rep(c, admin, masters):
    r = c.put("/api/grades/G1", json={"code": "G1", "name": "Sales executive", "da_in": 300, "km_in": 4, "hotel_in": 0,
                                      "da_out": 600, "km_out": 6, "hotel_out": 2000}, headers=H(admin))
    assert r.status_code == 200, r.text
    r = c.post("/api/users", json={"name": "Priya Verma", "username": "priya", "password": "priya12345", "role": "rep", "emp_code": "E1007",
                                   "grade": "G1"}, headers=H(admin))
    assert r.status_code == 200, r.text
    return login(c, "priya", "priya12345")


def cust_id(c, tok, code):
    return next(x["id"] for x in c.get("/api/customers", headers=H(tok)).json() if x["code"] == code)


def test_masters_are_read_only_and_downloaded(c, admin, masters):
    s = c.get("/api/integration/summary", headers=H(admin)).json()
    counts = {m["entity"]: m["count"] for m in s["masters"]}
    assert counts == {"customers": 5, "materials": 6, "attributes": 16, "prices": 11, "hsn": 4, "charges": 3}
    for path in ["/api/customers", "/api/materials", "/api/prices", "/api/hsn", "/api/charges"]:
        assert c.post(path, json={"code": "X", "name": "X"}, headers=H(admin)).status_code == 405
    tea = next(m for m in c.get("/api/materials", headers=H(admin)).json() if m["code"] == "SF-TEA-250")
    assert tea["gst"] == 5 and tea["attrs"]["Shelf life"] == "18 months" and tea["base_price"] == 210


def test_time_in_and_out(c, rep, admin):
    assert c.post("/api/timeclock/punch", json={"kind": "in"}, headers=H(rep)).status_code == 422          # location required
    r = c.post("/api/timeclock/punch", json={"kind": "in", "note": "Field", **LOC}, headers=H(rep))
    assert r.status_code == 200 and r.json()["session"]["in_note"] == "Field"
    assert c.post("/api/timeclock/punch", json={"kind": "in", **LOC}, headers=H(rep)).status_code == 409
    t = c.get("/api/timeclock", headers=H(rep)).json()
    assert t["on_duty"] and t["status"]["code"] == "P"


def test_day_status_late_short_and_correction(c, rep, admin):
    today = local_today()
    with SessionLocal() as db:
        u = db.query(User).filter_by(username="priya").one()
        u.joined_on = today - timedelta(days=30)
        def add(n, i, o):
            d = today - timedelta(days=n)
            db.add(Attendance(user_id=u.id, work_date=d, in_at=local_at(d, i), out_at=local_at(d, o) if o else None, in_lat=26.85, in_lng=80.95))
        add(1, "09:25", "18:10"); add(2, "10:05", "18:30"); add(3, "09:30", "12:15"); add(4, "09:40", None)
        db.commit()
    hist = {h["date"]: h for h in c.get("/api/timeclock", headers=H(rep)).json()["history"]}
    d = lambda n: (today - timedelta(days=n)).isoformat()  # noqa: E731
    assert hist[d(1)]["code"] == "P" and hist[d(1)]["late"] == 0
    assert hist[d(2)]["code"] == "P" and hist[d(2)]["late"] == 20
    assert hist[d(3)]["code"] == "S"
    assert c.get("/api/timeclock", headers=H(rep)).json()["stale"]["date"] == d(4)
    bad = c.post("/api/corrections", json={"date": d(4), "type": "out", "out_time": "08:00", "reason": "Phone battery died"}, headers=H(rep))
    assert bad.status_code == 422 and "after the Time In" in bad.json()["detail"]
    r = c.post("/api/corrections", json={"date": d(4), "type": "out", "out_time": "18:00", "reason": "Phone battery died"}, headers=H(rep))
    assert r.status_code == 200
    assert c.post(f"/api/corrections/{r.json()['id']}/decide", json={"approve": True}, headers=H(rep)).status_code == 403
    pend = c.get("/api/corrections", headers=H(admin)).json()
    assert len(pend) == 1
    assert c.post(f"/api/corrections/{pend[0]['id']}/decide", json={"approve": True}, headers=H(admin)).status_code == 200
    hist = {h["date"]: h for h in c.get("/api/timeclock", headers=H(rep)).json()["history"]}
    assert hist[d(4)]["code"] == "P" and hist[d(4)]["corrected"]
    mus = c.get(f"/api/attendance/muster?month={today.isoformat()[:7]}", headers=H(admin)).json()
    assert any(r["name"] == "Priya Verma" for r in mus["rows"])
    x = c.get(f"/api/attendance/muster.xlsx?month={today.isoformat()[:7]}", headers=H(admin))
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert c.get(f"/api/attendance/muster.xlsx?month={today.isoformat()[:7]}", headers=H(rep)).status_code == 403


def test_quotation_to_order_and_gst(c, rep, admin):
    awadh, patna = cust_id(c, rep, "C1001"), cust_id(c, rep, "C1004")
    base = {"kind": "quote", "customer_id": awadh, "date": local_today().isoformat(), "lines": [{"material": "SF-NDL-CS", "qty": 30}], **LOC}
    p = c.post("/api/docs/preview", json=base, headers=H(rep)).json()
    assert p["lines"][0]["price"] == 545 and p["lines"][0]["price_source"] == "DISTRIBUTOR price"   # 25+ cases
    assert p["totals"]["cgst"] == p["totals"]["sgst"] == 1471.5 and p["totals"]["total"] == 19293
    # a price sent from the device is ignored unless price changes are allowed
    p2 = c.post("/api/docs/preview", json={**base, "lines": [{"material": "SF-NDL-CS", "qty": 30, "price": 1}]}, headers=H(rep)).json()
    assert p2["lines"][0]["price"] == 545
    q = c.post("/api/docs", json=base, headers=H(rep)).json()
    assert q["status"] == "open" and q["no"].startswith("QT") and q["sync"]["status"] == "pending"
    o = c.post("/api/docs", json={**base, "kind": "order", "quote_id": q["id"], "charges": ["FRT"]}, headers=H(rep)).json()
    assert o["status"] == "submitted" and o["quote_no"] == q["no"] and o["sync"]["status"] == "na"
    t = o["totals"]
    assert t["pre"][0]["amount"] == 327 and t["taxable"] == 16677 and t["total"] == 19679   # 2% freight, GST 18%
    assert c.get(f"/api/docs/{q['id']}", headers=H(rep)).json()["status"] == "converted"
    assert c.post(f"/api/docs/{o['id']}/approve", headers=H(rep)).status_code == 403
    assert c.post(f"/api/docs/{o['id']}/approve", headers=H(admin)).json()["status"] == "approved"
    igst = c.post("/api/docs/preview", json={**base, "kind": "order", "customer_id": patna, "lines": [{"material": "SF-TEA-250", "qty": 20}]}, headers=H(rep)).json()
    assert igst["totals"]["igst"] == 189 and igst["totals"]["cgst"] == 0 and igst["totals"]["total"] == 3969
    blocked = cust_id(c, admin, "C1005")
    assert c.post("/api/docs", json={**base, "customer_id": blocked}, headers=H(admin)).status_code == 422


def test_payment_with_allocation(c, rep, admin):
    awadh = cust_id(c, rep, "C1001")
    oo = c.get(f"/api/customers/{awadh}/open-orders", headers=H(rep)).json()["orders"]
    assert len(oo) == 1 and oo[0]["due"] == 19679
    body = {"customer_id": awadh, "date": local_today().isoformat(), "amount": 5000, "mode": "neft", "ref": "UTR1", "allocations": [{"order_id": oo[0]["id"], "amount": 6000}], **LOC}
    assert c.post("/api/payments", json=body, headers=H(rep)).status_code == 422
    assert c.post("/api/payments", json={**body, "ref": ""}, headers=H(rep)).status_code == 422
    p = c.post("/api/payments", json={**body, "allocations": [{"order_id": oo[0]["id"], "amount": 5000}]}, headers=H(rep)).json()
    assert p["no"].startswith("RC") and p["allocations"][0]["amount"] == 5000
    assert c.get(f"/api/customers/{awadh}/open-orders", headers=H(rep)).json()["orders"][0]["due"] == 14679
    assert c.post(f"/api/docs/{oo[0]['id']}/cancel", headers=H(admin)).status_code == 409    # money allocated


def test_expense_claim(c, rep, admin):
    t = local_today()
    body = {"type": "out", "date_from": (t - timedelta(days=1)).isoformat(), "date_to": t.isoformat(), "purpose": "Distributor visits, Kanpur",
            "km": 120, "others": [{"category": "Food", "amount": 250, "note": "Lunch"}]}
    pv = c.post("/api/expenses/preview", json=body, headers=H(rep)).json()
    assert pv["nights"] == 1 and pv["hotel"] == 2000 and pv["calc"]["total"] == 4170
    x = c.post("/api/expenses", json={**body, "hotel": 2400}, headers=H(rep)).json()
    assert x["calc"]["total"] == 4570 and x["calc"]["hotel_over"] == 400 and x["status"] == "submitted"
    assert c.post("/api/expenses", json={**body, "type": "in"}, headers=H(rep)).status_code == 409        # overlapping days
    assert c.post(f"/api/expenses/{x['id']}/decide", json={"approve": True}, headers=H(rep)).status_code == 403
    r = c.post(f"/api/expenses/{x['id']}/decide", json={"approve": True}, headers=H(admin)).json()
    assert r["status"] == "approved" and r["sync"]["status"] == "pending"


def test_reps_see_only_their_own(c, admin, rep):
    r = c.post("/api/users", json={"name": "Second Rep", "username": "rep2", "password": "rep212345", "role": "rep"}, headers=H(admin))
    tok = login(c, "rep2", "rep212345")
    assert c.get("/api/docs?kind=order", headers=H(tok)).json() == []
    oid = c.get("/api/docs?kind=order", headers=H(rep)).json()[0]["id"]
    assert c.get(f"/api/docs/{oid}", headers=H(tok)).status_code == 403
    assert c.get("/api/integration/summary", headers=H(tok)).status_code == 403
    assert c.post("/api/users", json={"name": "x", "username": "x", "password": "x" * 8}, headers=H(tok)).status_code == 403


def test_dynamics_round_trip(c, admin):
    s = c.get("/api/integration/summary", headers=H(admin)).json()["out"]
    assert (s["quotes"]["count"], s["orders"]["count"], s["payments"]["count"], s["expenses"]["count"]) == (1, 1, 1, 1)
    r = c.post("/api/integration/export", headers=H(admin))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    rows = {n: wb[n].max_row - 1 for n in wb.sheetnames}
    assert rows["Quote Header"] == 1 and rows["Sales Header"] == 1 and rows["Payments"] == 1 and rows["Employee Expenses"] == 4
    pay = wb["Payments"]; hdr = [x.value for x in pay[1]]
    assert pay.cell(2, hdr.index("Amount") + 1).value == -5000
    nos = []
    for sheet, col in [("Quote Header", "External Document No."), ("Sales Header", "External Document No."), ("Payments", "Document No."), ("Employee Expenses", "Document No.")]:
        h = [x.value for x in wb[sheet][1]]
        nos += list({row[h.index(col)] for row in wb[sheet].iter_rows(min_row=2, values_only=True)})
    ack = Workbook(); ws = ack.active; ws.append(["External Document No.", "Document No."])
    for i, n in enumerate(nos):
        ws.append([n, f"DYN-{i}"])
    buf = io.BytesIO(); ack.save(buf)
    r = c.post("/api/integration/ack", files={"file": ("ack.xlsx", buf.getvalue())}, headers=H(admin)).json()
    assert r == {"matched": 4, "unknown": 0}
    s = c.get("/api/integration/summary", headers=H(admin)).json()
    assert s["waiting_file"] == 0 and all(v["count"] == 0 for v in s["out"].values())
    d = c.get("/api/dashboard", headers=H(admin)).json()
    assert d["staff"] == 2 and d["on_duty"] == 1 and d["collected"]["value"] == 5000
